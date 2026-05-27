import io
import os
import html as _html
import msal
import requests
import streamlit as st
import google.generativeai as genai
from docx import Document
from datetime import datetime, timedelta
from PIL import Image
import openpyxl
import re
from supabase import create_client

# --- 設定情報 ---
GEMINI_API_KEY = st.secrets["GEMINI_API_KEY"]
MS_CLIENT_ID = st.secrets["MS_CLIENT_ID"]
MS_TENANT_ID = st.secrets["MS_TENANT_ID"]
SUPABASE_URL = st.secrets["SUPABASE_URL"]
SUPABASE_KEY = st.secrets["SUPABASE_KEY"]

genai.configure(api_key=GEMINI_API_KEY)
supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

SCOPES = [
    "Files.Read.All",
    "Chat.Read",
    "Team.ReadBasic.All",
    "Channel.ReadBasic.All",
    "ChannelMessage.Read.All",
    "Notes.Read.All",
]

defaults = {
    "ms_token": None,
    "device_flow": None,
    "msal_app": None,
    "channels_list": None,
    "chat_history": [],
    "selected_channel_names": [],
}
for key, val in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = val

_avatar_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ブランキュ.png")
BRANCHU_AVATAR = Image.open(_avatar_path) if os.path.exists(_avatar_path) else "🐾"

SYSTEM_PROMPT = (
    "あなたは「ブランキュ」という名前の、福祉施設Plusらぼ専属のAIアシスタントです。\n"
    "元気で明るい性格で、スタッフを全力でサポートします。\n"
    "【口調のルール】\n"
    "・「〜だよ！」「〜してみるね！」「〜かな？」など親しみやすい口調で話す\n"
    "・情報を見つけたときは「あったよ！」「みつけた！」など元気に反応する\n"
    "・見つからないときも「ごめんね、見つからなかった…でも別のキーワードで試してみて！」など前向きに\n"
    "【回答のルール】\n"
    "・回答には必ず記録資料（いつ・誰が・どのOneNote/ファイル/メッセージ）のIDを含める\n"
    "・該当情報が複数あれば時系列で列挙する\n"
    "・福祉現場の用語や状況を幅広く理解して回答する\n"
    "・直接的な言葉がなくても文脈から関連すると判断できる情報も含める"
)


def get_msal_app():
    if st.session_state.msal_app is None:
        authority = "https://login.microsoftonline.com/" + MS_TENANT_ID
        st.session_state.msal_app = msal.PublicClientApplication(
            MS_CLIENT_ID, authority=authority
        )
    return st.session_state.msal_app


@st.cache_data(ttl=3600)
def get_generate_model_name():
    list_res = requests.get(
        "https://generativelanguage.googleapis.com/v1beta/models",
        params={"key": GEMINI_API_KEY}
    )
    if list_res.status_code == 200:
        models = list_res.json().get("models", [])
        for keyword in ["1.5-flash", "flash", "pro"]:
            for m in models:
                methods = m.get("supportedGenerationMethods", [])
                if "generateContent" in methods and keyword in m["name"]:
                    return m["name"].replace("models/", "")
    return None

def get_working_model():
    name = get_generate_model_name()
    if name:
        return genai.GenerativeModel(name)
    return genai.GenerativeModel("gemini-1.5-flash")


def graph_get(url, token):
    headers = {'Authorization': 'Bearer ' + token}
    res = requests.get(url, headers=headers)
    if res.status_code == 200:
        return res.json()
    return None


def strip_html(text):
    if not text:
        return ""
    return re.sub(r'<[^>]+>', '', text).strip()


def download_file_content(drive_id, item_id, token):
    url = (
        "https://graph.microsoft.com/v1.0/drives/"
        + drive_id + "/items/" + item_id + "/content"
    )
    headers = {'Authorization': 'Bearer ' + token}
    res = requests.get(url, headers=headers, allow_redirects=True)
    if res.status_code == 200:
        return res.content
    return None


def extract_text_from_bytes(file_bytes_raw, file_name):
    file_bytes = io.BytesIO(file_bytes_raw)
    text = ""
    try:
        if file_name.endswith(('.xlsx', '.xlsm')):
            wb = openpyxl.load_workbook(file_bytes, data_only=True)
            for sheet in wb.worksheets:
                text += "\n[シート: " + sheet.title + "]\n"
                for row in sheet.iter_rows(values_only=True):
                    row_data = " ".join([str(c) for c in row if c is not None])
                    if row_data.strip():
                        text += row_data + "\n"
        elif file_name.endswith('.docx'):
            doc = Document(file_bytes)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif file_name.endswith('.txt'):
            text = file_bytes_raw.decode('utf-8', errors='ignore')
        elif file_name.endswith('.pdf'):
            try:
                import fitz
                import base64
                pdf = fitz.open(stream=file_bytes_raw, filetype="pdf")
                for page in pdf:
                    img = page.get_pixmap(dpi=150)
                    b64 = base64.b64encode(img.tobytes("png")).decode()
                    vision_res = requests.post(
                        "https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent",
                        params={"key": GEMINI_API_KEY},
                        json={"contents": [{"parts": [
                            {"text": "この画像に書かれている文字をすべて読み取って、そのまま出力してください。"},
                            {"inline_data": {"mime_type": "image/png", "data": b64}}
                        ]}]}
                    )
                    if vision_res.status_code == 200:
                        text += vision_res.json()["candidates"][0]["content"]["parts"][0]["text"]
            except Exception as e:
                text = "(PDF解析エラー: " + str(e) + ")"
    except Exception as e:
        text = "(解析エラー: " + str(e) + ")"
    return text[:4000]


@st.cache_data(ttl=3600)
def get_embed_model_name():
    list_res = requests.get(
        "https://generativelanguage.googleapis.com/v1beta/models",
        params={"key": GEMINI_API_KEY}
    )
    if list_res.status_code == 200:
        models = list_res.json().get("models", [])
        for m in models:
            methods = m.get("supportedGenerationMethods", [])
            if "embedContent" in methods:
                return m["name"].replace("models/", "")
    return None

def get_embedding(text):
    import time
    model_name = get_embed_model_name()
    if not model_name:
        raise Exception("利用可能なembeddingモデルが見つかりません")
    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        + model_name + ":embedContent"
    )
    body = {
        "model": "models/" + model_name,
        "content": {"parts": [{"text": text[:2000]}]},
        "taskType": "RETRIEVAL_DOCUMENT"
    }
    for attempt in range(5):
        res = requests.post(url, json=body, params={"key": GEMINI_API_KEY})
        if res.status_code == 200:
            return res.json()["embedding"]["values"]
        if res.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        res.raise_for_status()
    res.raise_for_status()
    return []


# --- Supabaseにドキュメントを保存 ---
def save_document(source_type, source_id, title, content, author, recorded_at, url, channel_name, team_name):
    if not content or not content.strip():
        return
    try:
        embedding = get_embedding(content)
        existing = supabase.table("documents").select("id").eq("source_id", source_id).execute()
        if existing.data:
            supabase.table("documents").update({
                "content": content,
                "embedding": embedding,
                "updated_at": datetime.now().isoformat(),
            }).eq("source_id", source_id).execute()
        else:
            supabase.table("documents").insert({
                "source_type": source_type,
                "source_id": source_id,
                "title": title,
                "content": content,
                "embedding": embedding,
                "author": author,
                "recorded_at": recorded_at,
                "url": url,
                "channel_name": channel_name,
                "team_name": team_name,
            }).execute()
    except Exception as e:
        st.warning("DB保存エラー: " + str(e))


# --- Supabaseからベクトル検索 ---
def search_documents(query_text, channel_names=None):
    try:
        model_name = get_embed_model_name()
        if not model_name:
            st.warning("利用可能なembeddingモデルが見つかりませんでした。")
            return []
        embed_url = (
            "https://generativelanguage.googleapis.com/v1beta/models/"
            + model_name + ":embedContent"
        )
        embed_body = {
            "model": "models/" + model_name,
            "content": {"parts": [{"text": query_text}]},
            "taskType": "RETRIEVAL_QUERY"
        }
        embed_res = requests.post(embed_url, json=embed_body, params={"key": GEMINI_API_KEY})
        embed_res.raise_for_status()
        query_embedding = embed_res.json()["embedding"]["values"]
        result = supabase.rpc("match_documents", {
            "query_embedding": query_embedding,
            "match_threshold": 0.3,
            "match_count": 30,
            "filter_channels": channel_names or None,
        }).execute()
        return result.data or []
    except Exception as e:
        st.warning("DB検索エラー: " + str(e))
        return []


def get_teams_and_channels(token):
    items = []
    teams_data = graph_get("https://graph.microsoft.com/v1.0/me/joinedTeams", token)
    if teams_data:
        for team in teams_data.get('value', []):
            team_id = team['id']
            team_name = team['displayName']
            ch_url = "https://graph.microsoft.com/v1.0/teams/" + team_id + "/channels"
            ch_data = graph_get(ch_url, token)
            if ch_data:
                for ch in ch_data.get('value', []):
                    items.append({
                        'label': "📢 " + team_name + " / " + ch['displayName'],
                        'type': 'channel',
                        'team_id': team_id,
                        'team_name': team_name,
                        'channel_id': ch['id'],
                        'channel_name': ch['displayName'],
                    })
    chat_data = graph_get("https://graph.microsoft.com/v1.0/me/chats?$expand=members", token)
    if chat_data:
        for chat in chat_data.get('value', []):
            chat_id = chat['id']
            members = chat.get('members', [])
            names = [m.get('displayName', '') for m in members if m.get('displayName')]
            label = "、".join(names[:3]) if names else chat_id[:20]
            items.append({
                'label': "💬 " + label,
                'type': 'chat',
                'chat_id': chat_id,
            })
    return items


def index_channel(sel, token):
    team_id = sel['team_id']
    channel_id = sel['channel_id']
    team_name = sel.get('team_name', '')
    channel_name = sel.get('channel_name', '')
    count = 0

    one_year_ago = (datetime.utcnow() - timedelta(days=365)).strftime('%Y-%m-%dT%H:%M:%SZ')
    msg_url = (
        "https://graph.microsoft.com/v1.0/teams/"
        + team_id + "/channels/" + channel_id
        + "/messages?$top=50&$filter=createdDateTime ge " + one_year_ago
    )
    while msg_url:
        data = graph_get(msg_url, token)
        if not data:
            break
        msg_url = data.get('@odata.nextLink')
        for msg in data.get('value', []):
            body = strip_html(msg.get('body', {}).get('content', ''))
            sender = msg.get('from', {})
            user = sender.get('user', {}) if sender else {}
            name = user.get('displayName', '不明') if user else '不明'
            created = msg.get('createdDateTime', '')
            msg_id = msg.get('id', '')
            teams_link = (
                "https://teams.microsoft.com/l/message/"
                + channel_id + "/" + msg_id
                + "?groupId=" + team_id
                + "&tenantId=" + MS_TENANT_ID
            )
            atts = msg.get('attachments', [])
            att_names = [a.get('name', '') for a in atts if a.get('name')]
            full_content = body
            if att_names:
                full_content += " [添付: " + ", ".join(att_names) + "]"
            if full_content.strip():
                save_document(
                    'message', msg_id, None, full_content,
                    name, created, teams_link, channel_name, team_name
                )
                count += 1

            reply_url = (
                "https://graph.microsoft.com/v1.0/teams/"
                + team_id + "/channels/" + channel_id
                + "/messages/" + msg_id + "/replies?$top=20"
            )
            reply_data = graph_get(reply_url, token)
            if reply_data:
                for reply in reply_data.get('value', []):
                    rbody = strip_html(reply.get('body', {}).get('content', ''))
                    rsender = reply.get('from', {})
                    ruser = rsender.get('user', {}) if rsender else {}
                    rname = ruser.get('displayName', '不明') if ruser else '不明'
                    rcreated = reply.get('createdDateTime', '')
                    reply_id = reply.get('id', '')
                    reply_link = (
                        "https://teams.microsoft.com/l/message/"
                        + channel_id + "/" + reply_id
                        + "?groupId=" + team_id
                        + "&tenantId=" + MS_TENANT_ID
                    )
                    if rbody.strip():
                        save_document(
                            'message', reply_id, None, rbody,
                            rname, rcreated, reply_link, channel_name, team_name
                        )
                        count += 1

    def index_folder(drive_id, item_id):
        nonlocal count
        children_url = (
            "https://graph.microsoft.com/v1.0/drives/"
            + drive_id + "/items/" + item_id + "/children?$top=50"
        )
        while children_url:
            children_data = graph_get(children_url, token)
            if not children_data:
                break
            children_url = children_data.get('@odata.nextLink')
            for item in children_data.get('value', []):
                if 'folder' in item:
                    index_folder(drive_id, item['id'])
                    continue
                if 'file' not in item:
                    continue
                name = item['name']
                web_url = item.get('webUrl', '')
                file_item_id = item['id']
                file_drive_id = item.get('parentReference', {}).get('driveId', drive_id)
                supported = ('.xlsx', '.xlsm', '.docx', '.txt', '.pdf')
                if not name.endswith(supported):
                    continue
                content = download_file_content(file_drive_id, file_item_id, token)
                if content:
                    text = extract_text_from_bytes(content, name)
                    if text:
                        save_document(
                            'file', file_item_id, name, text,
                            None, None, web_url, channel_name, team_name
                        )
                        count += 1

    folder_url = (
        "https://graph.microsoft.com/v1.0/teams/"
        + team_id + "/channels/" + channel_id + "/filesFolder"
    )
    folder_data = graph_get(folder_url, token)
    if folder_data:
        drive_id = folder_data.get('parentReference', {}).get('driveId')
        item_id = folder_data.get('id')
        if drive_id and item_id:
            index_folder(drive_id, item_id)

    pages_url = (
        "https://graph.microsoft.com/v1.0/groups/"
        + team_id + "/onenote/pages"
        + "?$top=50&$select=id,title,createdDateTime,links"
    )
    pages_data = graph_get(pages_url, token)
    if pages_data:
        for page in pages_data.get('value', []):
            page_id = page.get('id', '')
            title = page.get('title', '無題')
            created = page.get('createdDateTime', '')
            page_links = page.get('links', {})
            one_note_url = page_links.get('oneNoteWebUrl', {}).get('href', '')
            content_url = (
                "https://graph.microsoft.com/v1.0/groups/"
                + team_id + "/onenote/pages/" + page_id + "/content"
            )
            headers = {'Authorization': 'Bearer ' + token}
            res = requests.get(content_url, headers=headers)
            if res.status_code == 200:
                body = strip_html(res.text)
                if body.strip():
                    save_document(
                        'onenote', page_id, title, body,
                        None, created, one_note_url, channel_name, team_name
                    )
                    count += 1

    return count


# ======================
# 吹き出し表示ヘルパー
# ======================
def _md_to_html(text):
    t = _html.escape(text)
    t = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', t)
    t = re.sub(r'^#{1,3}\s+(.+)$', r'<strong>\1</strong>', t, flags=re.MULTILINE)
    t = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2" target="_blank">\1</a>', t)
    t = t.replace('\n', '<br>')
    return t

def render_branchu_msg(content, links=None):
    col1, col2 = st.columns([1, 5])
    with col1:
        st.image(BRANCHU_AVATAR, width=100)
    with col2:
        st.markdown(
            '<div style="background:#f3e8ff; border-radius:20px; border-top-left-radius:4px;'
            ' padding:16px 20px; box-shadow:0 2px 6px rgba(0,0,0,0.08); line-height:1.7;">'
            + _md_to_html(content) + '</div>',
            unsafe_allow_html=True
        )
        if links:
            for link in links:
                if link.get("url"):
                    st.markdown("[" + link["label"] + "](" + link["url"] + ")")

def render_user_msg(content):
    col1, col2 = st.columns([5, 1])
    with col1:
        st.markdown(
            '<div style="background:#e8f4ff; border-radius:20px; border-top-right-radius:4px;'
            ' padding:16px 20px; line-height:1.7;">'
            + _md_to_html(content) + '</div>',
            unsafe_allow_html=True
        )
    with col2:
        st.markdown(
            '<div style="text-align:center; padding-top:10px; font-size:36px;">👤</div>',
            unsafe_allow_html=True
        )

# ======================
# UI
# ======================
st.title("ブランキュ AI検索アシスタント")

app = get_msal_app()

# --- ログインUI ---
if not st.session_state.ms_token:
    st.caption("メッセージ・ファイル・OneNote・PDFを横断検索するよ！まずはログインしてね！")
    if st.session_state.device_flow is None:
        if st.button("Microsoft 365 でログイン"):
            flow = app.initiate_device_flow(scopes=SCOPES)
            st.session_state.device_flow = flow
            st.rerun()
    else:
        flow = st.session_state.device_flow
        st.info(
            "以下のURLにアクセスしてコードを入力してください：\n\n"
            "**" + flow['verification_uri'] + "**\n\n"
            "コード：**" + flow['user_code'] + "**"
        )
        if st.button("ログイン完了（認証後にクリック）"):
            with st.spinner("認証確認中..."):
                result = app.acquire_token_by_device_flow(flow)
                if result and "access_token" in result:
                    st.session_state.ms_token = result["access_token"]
                    st.session_state.device_flow = None
                    st.rerun()
                else:
                    st.error("❌ 認証に失敗しました。")
                    st.session_state.device_flow = None

# --- ログイン後 ---
if st.session_state.ms_token:
    token = st.session_state.ms_token

    if st.session_state.channels_list is None:
        with st.spinner("Teams・チャット一覧を取得中..."):
            st.session_state.channels_list = get_teams_and_channels(token)

    channels = st.session_state.channels_list or []
    labels = [ch['label'] for ch in channels]

    # --- サイドバー ---
    with st.sidebar:
        st.success("✅ ログイン済み")
        if st.button("ログアウト"):
            for key in defaults:
                st.session_state[key] = defaults[key]
            st.rerun()

        st.divider()
        st.header("📂 検索先チャンネル")
        selected_indices = st.multiselect(
            "チャンネルを選んでね（複数OK）",
            range(len(labels)),
            format_func=lambda i: labels[i],
        )
        st.session_state.selected_channel_names = [
            channels[i].get('channel_name') for i in selected_indices
            if channels[i].get('channel_name')
        ]

        st.divider()
        st.header("🔄 インデックス更新")
        st.caption("新しい記録が増えたときに実行してね")
        index_indices = st.multiselect(
            "取り込むチャンネルを選んでね",
            range(len(labels)),
            format_func=lambda i: labels[i],
            key="index_select"
        )
        if st.button("🔄 インデックス更新を実行"):
            if not index_indices:
                st.warning("チャンネルを選んでください。")
            else:
                total_count = 0
                for sel_i in index_indices:
                    sel = channels[sel_i]
                    if sel['type'] != 'channel':
                        continue
                    with st.spinner(sel['label'] + " を取り込み中..."):
                        count = index_channel(sel, token)
                        total_count += count
                        st.write("✅ " + sel['label'] + ": " + str(count) + " 件")
                st.success("🎉 合計 " + str(total_count) + " 件保存したよ！")

    # --- 初回あいさつ ---
    if not st.session_state.chat_history:
        st.session_state.chat_history.append({
            "role": "assistant",
            "content": "こんにちは！ブランキュだよ！\nTeamsのメッセージやファイル、OneNoteを全力で探してくるね！\nまずはサイドバーで検索先チャンネルを選んでから、何でも聞いてね！",
            "links": [],
        })

    # --- 会話履歴表示 ---
    for msg in st.session_state.chat_history:
        if msg["role"] == "user":
            render_user_msg(msg["content"])
        else:
            render_branchu_msg(msg["content"], msg.get("links", []))

    # --- チャット入力 ---
    user_input = st.chat_input("ブランキュに聞いてみよう！例：Aさんの最近の体調は？")
    if user_input:
        if not st.session_state.selected_channel_names:
            st.warning("サイドバーで検索先チャンネルを選んでから聞いてね！")
        else:
            st.session_state.chat_history.append({"role": "user", "content": user_input, "links": []})
            render_user_msg(user_input)

            with st.spinner("探してるよ〜！"):
                all_docs = search_documents(user_input, st.session_state.selected_channel_names)

            if not all_docs:
                response = "ごめんね、DBにデータが見つからなかった…「インデックス更新」でデータを取り込んでみて！"
                render_branchu_msg(response)
                st.session_state.chat_history.append({"role": "assistant", "content": response, "links": []})
                else:
                    all_context = []
                    all_links = []
                    for doc in all_docs:
                        source_type = doc.get('source_type', '')
                        source_id = str(doc.get('source_id', '') or '')
                        title = str(doc.get('title', '') or '')
                        content = str(doc.get('content', '') or '')
                        author = str(doc.get('author', '不明') or '不明')
                        recorded_at = doc.get('recorded_at', '')
                        url = doc.get('url', '')

                        try:
                            dt = datetime.fromisoformat(recorded_at.replace('Z', '+00:00')) if recorded_at else None
                            date_str = dt.strftime('%Y/%m/%d %H:%M') if dt else ''
                        except Exception:
                            date_str = str(recorded_at or '')

                        if source_type == 'message':
                            entry = "[メッセージID:" + source_id + "] " + author + "（" + date_str + "）: " + content[:500]
                            icon, lbl = "📝", author + "（" + date_str + "）"
                        elif source_type == 'file':
                            entry = "[ファイルID:" + source_id + "] ファイル: " + title + ":\n" + content[:1000]
                            icon, lbl = "📄", title or source_id
                        else:
                            entry = "[OneNoteID:" + source_id + "] OneNote: " + title + "（" + date_str + "）:\n" + content[:2000]
                            icon, lbl = "📓", title + "（" + date_str + "）"

                        all_context.append(entry)
                        all_links.append({"id": source_id, "type": source_type, "label": icon + " " + lbl, "url": url})

                    context_text = "\n".join(all_context)
                    if len(context_text) > 50000:
                        context_text = context_text[:50000]

                    history_text = ""
                    for h in st.session_state.chat_history[-7:-1]:
                        role_label = "ユーザー" if h["role"] == "user" else "ブランキュ"
                        history_text += role_label + ": " + h["content"][:300] + "\n"

                    ai_prompt = (
                        SYSTEM_PROMPT + "\n\n"
                        + ("【これまでの会話】\n" + history_text + "\n" if history_text else "")
                        + "【今回の質問】\n" + user_input + "\n\n"
                        + "【関連データ】\n" + context_text
                    )

                    import time
                    response = ""
                    model = get_working_model()
                    with st.spinner("ブランキュが考えてるよ..."):
                        for attempt in range(5):
                            try:
                                ai_res = model.generate_content(ai_prompt)
                                response = ai_res.text.strip()
                                break
                            except Exception as e:
                                if "429" in str(e) and attempt < 4:
                                    time.sleep(2 ** attempt * 10)
                                    continue
                                response = "AI分析エラー: " + str(e)
                                break

                    shown_links = [lnk for lnk in all_links if lnk["id"] in response and lnk["url"]]
                    display_links = shown_links if shown_links else [lnk for lnk in all_links[:20] if lnk["url"]]
                    render_branchu_msg(response, display_links)

                    st.session_state.chat_history.append({
                        "role": "assistant",
                        "content": response,
                        "links": display_links,
                    })
